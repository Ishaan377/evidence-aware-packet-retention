"""Create the all-in-one guide from authored content and executed result files."""
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
content=json.loads((ROOT/"docs/beginner_guide_content.json").read_text(encoding="utf-8"))
sources=json.loads((ROOT/"docs/sources.json").read_text(encoding="utf-8"))
summary=json.loads((ROOT/"results/demo/aggregate_metrics.json").read_text(encoding="utf-8"))
answers=json.loads((ROOT/"results/demo/representative_answers.json").read_text(encoding="utf-8"))
manifest=json.loads((ROOT/"results/demo/manifest.json").read_text(encoding="utf-8"))
for name,expected in manifest["source_sha256"].items():
    assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==expected,name

doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5)
section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.65)
section.left_margin=section.right_margin=Inches(.75)
section.footer_distance=Inches(.25)
for name,size in [("Normal",11),("Title",25),("Subtitle",12),("Heading 1",17),("Heading 2",12.5)]:
    style=doc.styles[name]
    style.font.name="Arial"
    style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.underline=False
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.1
    if name.startswith("Heading"):
        style.font.bold=True
        style.paragraph_format.keep_with_next=True
        style.paragraph_format.space_before=Pt(8 if name=="Heading 2" else 0)
    props=style.element.find(qn("w:pPr"))
    if props is not None:
        borders=props.find(qn("w:pBdr"))
        if borders is not None:
            props.remove(borders)
doc.styles["Normal"].paragraph_format.widow_control=True
doc.core_properties.title=content["title"]
doc.core_properties.subject="Project fundamentals professor demonstration and final research plan"
doc.core_properties.author=""
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run("Page ")
run.font.name="Arial"
run.font.size=Pt(9)
field=OxmlElement("w:fldSimple")
field.set(qn("w:instr"),"PAGE")
footer._p.append(field)

def para(text,style=None):
    p=doc.add_paragraph(text,style=style)
    p.paragraph_format.space_after=Pt(6)
    return p

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    props=t._tbl.tblPr
    borders=OxmlElement("w:tblBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        b=OxmlElement("w:"+side)
        b.set(qn("w:val"),"single")
        b.set(qn("w:sz"),"4")
        b.set(qn("w:color"),"D9D9D9")
        borders.append(b)
    props.append(borders)
    for col,width in zip(t.columns,widths):
        col.width=Inches(width)
    allrows=[headers]+rows
    for i,values in enumerate(allrows):
        row=t.rows[0] if i==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i==0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j,(cell,value,width) in enumerate(zip(row.cells,values,widths)):
            cell.width=Inches(width)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement("w:tcMar")
            for side in ("top","left","bottom","right"):
                el=OxmlElement("w:"+side)
                el.set(qn("w:w"),"95")
                el.set(qn("w:type"),"dxa")
                margins.append(el)
            tcpr.append(margins)
            shd=OxmlElement("w:shd")
            shd.set(qn("w:fill"),"DCE6EF" if i==0 else ("F5F7F9" if i%2==0 else "FFFFFF"))
            tcpr.append(shd)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(1)
            p.paragraph_format.line_spacing=1.03
            p.paragraph_format.keep_with_next=False
            if widths[j]<1.1 or (i==0 and widths[j]<2):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(str(value))
            r.font.name="Arial"
            r.font.size=Pt(10)
            r.font.color.rgb=RGBColor(0,0,0)
            r.bold=i==0
    gap=doc.add_paragraph()
    gap.paragraph_format.space_after=Pt(1)
    gap.paragraph_format.space_before=Pt(0)
    gap.paragraph_format.line_spacing=1
    gap.add_run().font.size=Pt(2)
    return t

def hyperlink(p,text,url):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement("w:hyperlink")
    h.set(qn("r:id"),rel)
    r=OxmlElement("w:r")
    rp=OxmlElement("w:rPr")
    color=OxmlElement("w:color")
    color.set(qn("w:val"),"000000")
    rp.append(color)
    u=OxmlElement("w:u")
    u.set(qn("w:val"),"single")
    rp.append(u)
    size=OxmlElement("w:sz")
    size.set(qn("w:val"),"20")
    rp.append(size)
    r.append(rp)
    tx=OxmlElement("w:t")
    tx.text=text
    r.append(tx)
    h.append(r)
    p._p.append(h)

def result_table():
    rows=[]
    for b in (.01,.05,.10,.20,1):
        group={r["strategy"]:r for r in summary if r["budget_fraction"]==b}
        random=group["random"]
        rows.append([f"{b*100:g}%",f'{random["mean_coverage"]*100:.1f}% +/- {random["sd_coverage"]*100:.2f} pp',
                     *[f'{group[s]["mean_coverage"]*100:g}%' for s in ("alert","evidence","flow_prefix")]])
    table(["Cap","Random mean and SD","Alert","Evidence","Prefix"],rows,[.65,2.3,1.1,1.5,1.2])

def matrix():
    names=["Q1 scan identity","Q2 exact count","Q3 interval","Q4 rejection pairs",
           "Q5 HTTP 200 pair","Q6 handshake","Q7 DNS pair","Q8 byte threshold"]
    groups=[answers[f"{s}_b0.05_seed0"]["questions"] for s in ("evidence","alert","random","flow_prefix")]
    rows=[[name]+["YES" if g[i]["answerable"] else "NO" for g in groups] for i,name in enumerate(names)]
    table(["Criterion at 5%","Evidence","Alert","Random seed 0","Prefix"],rows,[2.2,1.05,.8,1.9,.8])
    para("This matrix shows representative seed-zero files. It is not the random mean across thirty trials.")

for i,page in enumerate(content["pages"]):
    p=doc.add_paragraph(page["title"],style="Title" if page.get("cover") else "Heading 1")
    if i:
        p.paragraph_format.page_break_before=True
    p.paragraph_format.space_after=Pt(10)
    for block in page["blocks"]:
        kind=block["type"]
        if kind=="p":
            para(block["text"])
        elif kind=="h2":
            para(block["text"],"Heading 2")
        elif kind=="code":
            p=doc.add_paragraph()
            p.paragraph_format.space_before=Pt(2)
            p.paragraph_format.space_after=Pt(7)
            p.paragraph_format.keep_together=True
            for j,line in enumerate(block["text"].split("\n")):
                if j:
                    p.add_run().add_break()
                r=p.add_run(line)
                r.font.name="Consolas"
                r.font.size=Pt(9.5)
        elif kind=="table":
            table(block["headers"],block["rows"],block["widths"])
        elif kind=="bullets":
            for item in block["items"]:
                p=para(item,"List Bullet")
                p.paragraph_format.space_after=Pt(5)
        elif kind=="image":
            path=ROOT/block["path"]
            with Image.open(path) as img:
                iw,ih=img.size
            width=min(block["width"],block["max_height"]*iw/ih)
            p=doc.add_paragraph()
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            picture=p.add_run().add_picture(str(path),width=Inches(width))
            picture._inline.docPr.set("descr",block["caption"])
            cap=para(block["caption"])
            cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after=Pt(8)
            for r in cap.runs:
                r.font.size=Pt(9)
                r.italic=True
        elif kind=="results_table":
            result_table()
        elif kind=="question_matrix":
            matrix()
        elif kind=="source_links":
            for sid in block["ids"]:
                source=next(s for s in sources if s["id"]==sid)
                p=doc.add_paragraph()
                p.paragraph_format.space_after=Pt(4)
                year=source["year"] if source["year"] else "undated"
                label=block.get("labels",{}).get(str(sid),source["title"])
                hyperlink(p,f'[{sid}] {label} ({year})',source["url"])
        else:
            raise ValueError(kind)
out=ROOT/"docs/zero_to_hero_project_guide.docx"
doc.save(out)
print("Created",out,"with",len(content["pages"]),"planned pages.")
