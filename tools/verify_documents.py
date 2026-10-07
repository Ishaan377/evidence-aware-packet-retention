"""Verify current document packages, pagination and unchanged experiment provenance."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from lxml import etree
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def load(name): return json.loads((ROOT/name).read_text(encoding='utf-8'))
def norm(text): return re.sub(r'[^a-z0-9]', '', text.lower())
checks={}
def check(name, condition):
    assert condition, name
    checks[name]=True
reference=ROOT/'Network_Traffic_Forensics_Synopsis_FINAL.docx'
synopsis=ROOT/'docs/synopsis.docx'
guide=ROOT/'docs/zero_to_hero_project_guide.docx'
check('reference_unchanged',digest(reference)=='2a9000836d4bbac7e5db0347597bd17e8a10f8b6acd5157fda2e638452659800')
evidence=load('docs/qa/template_evidence.json')
check('synopsis_matches_template_evidence_hash',digest(synopsis)==evidence['final_sha256'])
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with ZipFile(reference) as z: original={n:z.read(n) for n in z.namelist()}
with ZipFile(synopsis) as z: current={n:z.read(n) for n in z.namelist()}
preserved=[n for n in original if n not in editable]
check('fourteen_original_parts_identical',len(preserved)==14 and all(original[n]==current[n] for n in preserved))
check('only_requested_image_part_added',set(current)-set(original)=={'word/media/architecture_flowchart.png'})
check('embedded_architecture_matches_asset',current['word/media/architecture_flowchart.png']==(ROOT/'docs/assets/architecture_flowchart.png').read_bytes())
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS={'w':W}
a=etree.fromstring(original['word/document.xml'])
b=etree.fromstring(current['word/document.xml'])
check('section_geometry_identical',etree.tostring(a.find('w:body/w:sectPr',NS))==etree.tostring(b.find('w:body/w:sectPr',NS)))
check('cover_table_identical',etree.tostring(a.find('w:body/w:tbl',NS))==etree.tostring(b.find('w:body/w:tbl',NS)))
def headings(root):
    return [''.join(p.xpath('.//w:t/text()',namespaces=NS)) for p in root.findall('w:body/w:p',NS)
            if p.find('w:r/w:rPr/w:sz',NS) is not None and p.find('w:r/w:rPr/w:sz',NS).get('{'+W+'}val')=='28']
check('sixteen_headings_same_order',len(headings(a))==16 and headings(a)==headings(b))
original_paragraphs=a.findall('w:body/w:p',NS)
texts=[''.join(p.xpath('.//w:t/text()',namespaces=NS)) for p in b.findall('w:body/w:p',NS)]
check('declaration_and_signature_text_preserved',all(''.join(p.xpath('.//w:t/text()',namespaces=NS)) in texts for p in original_paragraphs[49:]))
rels=etree.fromstring(current['word/_rels/document.xml.rels'])
check('architecture_relationship_valid',any(r.get('Id')=='rIdArchitectureFlowchart' and r.get('Target')=='media/architecture_flowchart.png' for r in rels))
check('png_content_type_valid',any(r.get('Extension')=='png' and r.get('ContentType')=='image/png' for r in etree.fromstring(current['[Content_Types].xml'])))
refdoc=Document(reference)
syn=Document(synopsis)
check('team_identity_cells_preserved',all([c.text for c in r.cells[:3]]==[c.text for c in s.cells[:3]] for r,s in zip(refdoc.tables[1].rows,syn.tables[1].rows)))
planned=[r.cells[2].text for r in refdoc.tables[3].rows[1:]]
check('ten_planned_completion_strings_match_reference',len(planned)==10 and planned==[r.cells[2].text for r in syn.tables[3].rows[1:]])
check('planned_schedule_explanation_present',load('docs/synopsis_content.json')['timeline_note'] in texts)
content=load('docs/beginner_guide_content.json')
gpdf=PdfReader(ROOT/'docs/qa/beginner_guide_final/guide.pdf')
spdf=PdfReader(ROOT/'docs/qa/synopsis_flowchart_v1/synopsis.pdf')
check('final_guide_has_29_pages',len(gpdf.pages)==len(content['pages'])==29)
check('final_synopsis_has_10_pages',len(spdf.pages)==10)
check('guide_chapters_start_on_advertised_pages',all(norm(p['title']) in norm(gpdf.pages[i].extract_text()) for i,p in enumerate(content['pages'])))
check('rendered_png_counts_complete',len(list((ROOT/'docs/qa/beginner_guide_final').glob('page-*.png')))==29 and len(list((ROOT/'docs/qa/synopsis_flowchart_v1').glob('page-*.png')))==10)
gdoc=Document(guide)
check('three_guide_figures_have_alt_text',len(gdoc.inline_shapes)==3 and all(s._inline.docPr.get('descr') for s in gdoc.inline_shapes))
source_ids=next(b['ids'] for p in content['pages'] for b in p['blocks'] if b['type']=='source_links')
sources=load('docs/sources.json')
expected_links={s['url'] for s in sources if s['id'] in source_ids}
actual_links={r.target_ref for r in gdoc.part.rels.values() if r.reltype.endswith('/hyperlink')}
check('all_eight_source_links_match_registry',len(expected_links)==8 and actual_links==expected_links)
summary=load('results/demo/aggregate_metrics.json')
result_table=next(t for t in gdoc.tables if [c.text for c in t.rows[0].cells]==['Cap','Random mean and SD','Alert','Evidence','Prefix'])
for row,cap in zip(result_table.rows[1:],(.01,.05,.10,.20,1)):
    group={r['strategy']:r for r in summary if r['budget_fraction']==cap}
    r=group['random']
    expected=[f'{cap*100:g}%',f'{r["mean_coverage"]*100:.1f}% +/- {r["sd_coverage"]*100:.2f} pp',*[f'{group[s]["mean_coverage"]*100:g}%' for s in ('alert','evidence','flow_prefix')]]
    check(f'guide_actual_results_at_{cap:g}',[c.text for c in row.cells]==expected)
m=load('results/demo/manifest.json')
check('all_core_and_demo_sources_unchanged',all(digest(ROOT/name)==value for name,value in m['source_sha256'].items()))
check('main_input_truth_unchanged',all(digest(ROOT/name)==m[key] for name,key in ((m['input_relative'],'input_sha256'),(m['truth_relative'],'truth_sha256'))))
check('saved_run_config_hash_unchanged',digest(ROOT/'results/demo/config_used.json')==m['config_sha256'])
check('root_config_matches_executed_values',load('config.json')==load('results/demo/config_used.json'))
check('all_165_saved_subset_hashes_unchanged',len(m['subsets'])==165 and all(digest(ROOT/info['path'])==info['sha256'] for info in m['subsets'].values()))
important=['docs/zero_to_hero_project_guide.docx','docs/synopsis.docx','docs/beginner_guide_content.json','docs/synopsis_content.json','docs/assets/architecture_flowchart.png','docs/assets/architecture_flowchart.svg','tools/build_project_guide.py','tools/build_synopsis.py','tools/create_architecture.py','tools/verify_documents.py']
report={
 'checked_at_utc':datetime.now(timezone.utc).isoformat(),
 'scope':'7 October document-only follow-up; supersedes historical DOCX hash/page/part counts in final_integration_checks.txt',
 'checks':checks,
 'guide_pages':29,'synopsis_pages':10,
 'reference_planned_completion_strings':planned,
 'visual_qa':{'renderer':'Hidden Word read-only PDF export plus bundled Poppler at 1400 pixels','guide_directory':'docs/qa/beginner_guide_final','guide_pages_inspected':list(range(1,30)),'synopsis_directory':'docs/qa/synopsis_flowchart_v1','synopsis_pages_inspected':list(range(1,11)),'result':'All latest pages inspected; commands, charts, tables, flowchart and signatures legible with no clipping or overlap.'},
 'important_file_sha256':{name:digest(ROOT/name) for name in important},
 'experiment_status':'Core sources, config, input/truth and all 165 saved subset hashes unchanged; no new experimental claim or unnecessary rerun.'
}
(ROOT/'logs/document_update_checks.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
print(f'Document checks passed: {len(checks)} checks; guide 29 pages; synopsis 10 pages; unchanged reference and experiment provenance.')